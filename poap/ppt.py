from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
import pandas as pd

NAVY=RGBColor(20,48,73); BLUE=RGBColor(28,100,160); PALE=RGBColor(245,248,251); LINE=RGBColor(210,218,226); GREY=RGBColor(105,115,125)
RAG={'GREEN':RGBColor(54,142,74),'AMBER':RGBColor(224,155,32),'RED':RGBColor(196,62,62)}

def _text(slide,x,y,w,h,text,size=10,bold=False,color=NAVY,align=None):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=s.text_frame; tf.clear(); tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=str(text); p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
    if align is not None: p.alignment=align
    return s

def _box(slide,x,y,w,h,title,body=''):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=PALE; s.line.color.rgb=LINE
    _text(slide,x+.16,y+.10,w-.32,.28,title,11,True,BLUE)
    if body: _text(slide,x+.16,y+.44,w-.32,h-.52,body,9,False,NAVY)
    return s

def _metric(slide,x,y,w,label,value):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.72)); s.fill.solid(); s.fill.fore_color.rgb=PALE; s.line.color.rgb=LINE
    _text(slide,x+.12,y+.08,w-.24,.22,label.upper(),7,True,BLUE); _text(slide,x+.12,y+.30,w-.24,.30,value,15,True,NAVY)

def _dt(v):
    if v is None or str(v) in ('','None','NaT'): return None
    try: return pd.to_datetime(v)
    except Exception: return None

def _date(v):
    d=_dt(v); return d.strftime('%d %b %Y') if d is not None else 'N/A'

def _milestone_data(summary):
    raw=summary.get('milestones',[]); result=[]
    if hasattr(raw,'iterrows'):
        for _,r in raw.iterrows():
            d=r.get('finish') if pd.notna(r.get('finish')) else r.get('start'); result.append({'task':str(r.get('task','')),'date':_dt(d),'rag':str(r.get('rag','GREEN')).upper()})
    else:
        for r in raw: result.append({'task':str(r.get('task','')),'date':_dt(r.get('date')),'rag':str(r.get('rag','GREEN')).upper()})
    return [m for m in result if m['date'] is not None]

def _timeline(slide,summary,x=.45,y=6.02,w=12.43,h=.88):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=PALE; s.line.color.rgb=LINE
    _text(slide,x+.16,y+.08,1.0,.2,'TIMELINE',9,True,BLUE)
    start=_dt(summary.get('start')); finish=_dt(summary.get('finish')); milestones=_milestone_data(summary)
    if start is None or finish is None or finish<=start:
        _text(slide,x+1.15,y+.08,w-1.35,.25,f"{_date(start)}  —  {_date(finish)}",8,False,GREY); return
    left=x+.35; right=x+w-.35; line_y=y+.49; span=(finish-start).total_seconds()
    line=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(left),Inches(line_y),Inches(right-left),Inches(.025)); line.fill.solid(); line.fill.fore_color.rgb=LINE; line.line.fill.background()
    # Month labels, capped for readability.
    months=pd.date_range(start=start.replace(day=1),end=finish,freq='MS')
    step=max(1,int(len(months)/7)+1)
    for d in months[::step]:
        pos=max(0,min(1,(d-start).total_seconds()/span)); px=left+(right-left)*pos
        tick=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(px),Inches(line_y-.05),Inches(.012),Inches(.12)); tick.fill.solid(); tick.fill.fore_color.rgb=GREY; tick.line.fill.background(); _text(slide,px-.25,y+.64,.55,.15,d.strftime('%b %y'),6,False,GREY,PP_ALIGN.CENTER)
    # Current date marker when inside plan range.
    today=pd.Timestamp.today().normalize()
    if start<=today<=finish:
        pos=(today-start).total_seconds()/span; px=left+(right-left)*pos
        now=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(px),Inches(y+.27),Inches(.018),Inches(.39)); now.fill.solid(); now.fill.fore_color.rgb=BLUE; now.line.fill.background(); _text(slide,px-.22,y+.15,.46,.12,'TODAY',5,True,BLUE,PP_ALIGN.CENTER)
    # Milestone diamonds with abbreviated labels above/below the line.
    for i,m in enumerate(milestones[:10]):
        pos=max(0,min(1,(m['date']-start).total_seconds()/span)); px=left+(right-left)*pos; color=RAG.get(m['rag'],BLUE)
        diamond=slide.shapes.add_shape(MSO_SHAPE.DIAMOND,Inches(px-.055),Inches(line_y-.055),Inches(.11),Inches(.11)); diamond.fill.solid(); diamond.fill.fore_color.rgb=color; diamond.line.fill.background()
        label=m['task'][:22]+('…' if len(m['task'])>22 else ''); ly=y+.24 if i%2==0 else y+.52
        _text(slide,max(x+.1,min(px-.48,x+w-1.05)),ly,.96,.18,label,5,False,NAVY,PP_ALIGN.CENTER)

def generate(summary,out_path):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); slide=prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide,.45,.25,10.6,.45,summary['project'],24,True,NAVY); _text(slide,.45,.73,8,.25,'PLAN ON A PAGE • EXECUTIVE VIEW',9,True,BLUE)
    rag=str(summary.get('overall_rag',summary.get('status','GREEN'))).upper(); rag=rag if rag in RAG else 'GREEN'
    badge=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(11.55),Inches(.28),Inches(1.3),Inches(.55)); badge.fill.solid(); badge.fill.fore_color.rgb=RAG[rag]; badge.line.fill.background(); _text(slide,11.55,.41,1.3,.22,rag,10,True,RGBColor(255,255,255),PP_ALIGN.CENTER)
    sd=_date(summary.get('start')); ed=_date(summary.get('finish')); tasks=summary.get('task_count',summary.get('tasks',0)); complete=summary.get('average_complete',summary.get('completion',0)); deps=summary.get('dependencies',0); rc=summary.get('rag_counts',{})
    _metric(slide,.45,1.15,2.0,'Tasks',str(tasks)); _metric(slide,2.58,1.15,2.0,'Average complete',f"{complete}%"); _metric(slide,4.71,1.15,2.25,'Plan start',sd); _metric(slide,7.09,1.15,2.25,'Plan finish',ed); _metric(slide,9.47,1.15,1.65,'Dependencies',str(deps)); _metric(slide,11.25,1.15,1.6,'R / A / G',f"{rc.get('RED',0)} / {rc.get('AMBER',0)} / {rc.get('GREEN',0)}")
    _box(slide,.45,2.10,3.15,2.08,'WORKSTREAMS','\n'.join('• '+x for x in summary.get('workstreams',[])))
    md=_milestone_data(summary); _box(slide,3.75,2.10,5.65,2.08,'KEY MILESTONES','\n'.join(f"• {m['date'].strftime('%d %b %y')}   {m['task']}" for m in md[:8]) or 'No dated milestones detected')
    _box(slide,9.55,2.10,3.33,2.08,'STATUS','\n'.join([f"• Green: {rc.get('GREEN',0)}",f"• Amber: {rc.get('AMBER',0)}",f"• Red: {rc.get('RED',0)}",f"• Overall: {rag}"]))
    _box(slide,.45,4.36,6.15,1.48,'RISKS / ATTENTION','\n'.join('• '+x for x in summary.get('risks',[])) or '• No rule-based risks detected')
    executive=summary.get('executive_summary') or f"{tasks} activities across {len(summary.get('workstreams',[]))} detected workstream(s). Average recorded completion is {complete}%."; _box(slide,6.75,4.36,6.13,1.48,'EXECUTIVE SUMMARY',executive)
    _timeline(slide,summary)
    _text(slide,.45,7.08,12.4,.18,'Generated by POAP Builder • All PowerPoint shapes and text remain editable',7,False,GREY,PP_ALIGN.RIGHT); prs.save(out_path)